import csv
import openpyxl
import os

def export_questions():
    excel_path = r'C:\Users\何任軒\Desktop\待整理\星墜答問 題庫與紀錄.xlsx'
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Master Excel file not found: {excel_path}")
    
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    # Support sheet name variations
    sheet_name = None
    for name in ['Questions', 'Question', 'Quetions', '題庫']:
        if name in wb.sheetnames:
            sheet_name = name
            break
    if not sheet_name:
        raise ValueError(f"Could not find Questions sheet in {wb.sheetnames}")
    
    ws = wb[sheet_name]
    output_csv = os.path.join(os.path.dirname(__file__), '..', 'data', 'default-question-bank.csv')
    output_csv = os.path.abspath(output_csv)
    
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    # Filter trailing empty headers if any
    last_col = len(headers)
    while last_col > 0 and not headers[last_col - 1]:
        last_col -= 1
    headers = headers[:last_col]
    
    rows = []
    for r in range(2, ws.max_row + 1):
        # Check if row is empty
        row_vals = []
        is_empty = True
        for col in range(1, last_col + 1):
            val = ws.cell(r, col).value
            if val is not None and str(val).strip() != '':
                is_empty = False
            
            # Formatting
            if isinstance(val, bool):
                val_str = 'TRUE' if val else 'FALSE'
            elif isinstance(val, float) and val.is_integer():
                val_str = str(int(val))
            elif val is None:
                val_str = ''
            else:
                val_str = str(val).strip()
            row_vals.append(val_str)
            
        if not is_empty and row_vals[0]: # Must have question_id
            rows.append(row_vals)
            
    print(f"Exporting {len(rows)} questions from '{sheet_name}' to '{output_csv}'...")
    with open(output_csv, 'w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(headers)
        writer.writerows(rows)
        
    print(f"Successfully exported {len(rows)} questions to {output_csv}")
    return len(rows)

if __name__ == '__main__':
    export_questions()
